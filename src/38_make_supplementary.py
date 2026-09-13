# -*- coding: utf-8 -*-
"""Build supplementary.docx: one page per supplementary figure (S1-S11) and
real Word tables for Tables S1-S12, from the pipeline result CSVs.

Captions for the tables are parsed live from manuscript/submission/
manuscript.md so there is a single source of truth.
"""
import csv
import io
import re
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
STAT = ROOT / "results" / "statistics"
OUT = ROOT / "manuscript" / "submission"

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def fmt(v):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if abs(x - round(x)) < 1e-12 and abs(x) < 1e6:
        return f"{int(round(x)):,}" if abs(x) >= 1000 else str(int(round(x)))
    return f"{x:.3f}"


def rows_of(fname, cols=None, keep=None):
    with open(STAT / fname, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if keep:
        rows = [r for r in rows if keep(r)]
    if cols:
        rows = [{c: r[c] for c in cols} for r in rows]
    return rows


FIG_S = {
    1: "Figure S1. Assay-level distribution of the Spearman correlation between "
       "disagreement and ensemble prediction error (median 0.032; Section 2.6).",
    2: "Figure S2. Mean ensemble prediction error by disagreement decile "
       "(monotonic: top decile 0.214 vs bottom decile 0.201).",
    3: "Figure S3. Disagreement versus distance to the nearest UniProt domain "
       "boundary (no boundary effect).",
    4: "Figure S4. BioGate 5-fold UniProt-grouped cross-validation: protein-level "
       "Spearman distributions against all baselines.",
    5: "Figure S5. BioGate context-feature ablation: validation Spearman for the "
       "full gate versus removal of pLDDT (no_plddt), functional annotations "
       "(no_ann) or MSA-depth (no_msa) context features. The "
       "scores-only / context-only / scores+context XGBoost comparison is in Table S5.",
    6: "Figure S6. Leave-one-model-out sensitivity of the disagreement-error "
       "association (range 0.027-0.039).",
    7: "Figure S7. Percentile (u) versus inverse-normal (z) disagreement scale: "
       "median assay-level rho 0.145 vs 0.032 at fixed error definition.",
    8: "Figure S8. Disagreement-error association by assay selection type "
       "(stability 0.076, binding 0.057; activity 0.003, organismal fitness -0.004).",
    9: "Figure S9. Residue mechanics (secondary structure, burial, contact "
       "density) versus family advantage or disagreement: null result.",
    10: "Figure S10. ESM-IF1 gain from experimental structures over AlphaFold2 "
        "by assay mean pLDDT (64 assays; low-confidence assays show MORE negative "
        "gain: -0.081 vs -0.023, difference -0.058, 95% CI -0.099 to -0.013).",
    11: "Figure S11. SSEmb distinct-architecture replication: residue-level "
        "Spearman between pLDDT and evolution-vs-SSEmb disagreement (median -0.067; "
        "65.6% of 151 proteins negative).",
}

md = (OUT / "manuscript.md").read_text(encoding="utf-8")
seg = md[md.find("**Table S1."):]
seg = seg[:seg.find("## Figure captions") if "## Figure captions" in seg[:200] else len(seg)]
TBL_CAP = {}
parts = re.split(r"(?=\*\*Table S\d)", seg)
for pt in parts:
    m = re.match(r"\*\*Table (S\d+)\.(.*?)\n\n", pt, re.S)
    if not m:
        m = re.match(r"\*\*Table (S\d+)\.(.*)$", pt, re.S)
    if m:
        TBL_CAP[m.group(1)] = ("Table " + m.group(1) + "." + m.group(2)).replace("**", "")
assert len(TBL_CAP) == 12, sorted(TBL_CAP)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(10.5)


def para(text, bold=False, size=None, center=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_table(headers, data, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, headers):
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
    for row in data:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            r = c.paragraphs[0].add_run(str(v))
            r.font.size = Pt(9)
    return t


# ---- cover page ----
para("Supplementary Information", bold=True, size=16, center=True)
para("Model Disagreement as a Scientific Observable: Structure and Boundaries of Protein AI Divergence",
     bold=True, size=13, center=True)
para("")
para("Contents", bold=True)
para("Supplementary Figures S1-S11 (one figure per page).", size=10.5)
para("Supplementary Tables S1-S12.", size=10.5)
para("")
para("All underlying machine-readable tables and pipeline code (scripts 01-44) "
     "are archived with the manuscript at "
     "https://github.com/hmjpan/protein-ai-disagreement (results/ and figures/ "
     "directories). Primary inferential summaries aggregate at the protein "
     "level where applicable, with bootstrap inference used for the principal "
     "effect estimates; assay-, residue- and variant-level quantities are "
     "labelled as such at each point of use. The analysis plan and GO/NO-GO "
     "thresholds were fixed in config.yaml before the final analyses.", size=10.5)

# ---- supplementary figures ----
doc.add_page_break()
para("Supplementary Figures", bold=True, size=14)
figdir = ROOT / "figures" / "supplementary"
for n in range(1, 12):
    files = sorted(f for f in figdir.iterdir()
                   if re.match(rf"FigS{n}[_\.]", f.name))
    doc.add_page_break()
    para(FIG_S[n], bold=True, size=10.5)
    for f in files:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(f), width=Inches(6.2))

# ---- supplementary tables ----
doc.add_page_break()
para("Supplementary Tables", bold=True, size=14)


def cap(key):
    doc.add_page_break()
    para(TBL_CAP[key].replace("\n", " "), bold=True, size=10.5)


# S1: disagreement definitions
cap("S1")
add_table(["Definition", "n proteins", "median rho", "95% CI",
           "frac proteins negative"],
          [[r["definition"], r["n_proteins"], fmt(r["median_rho"]),
            f"{fmt(r['ci_low'])} to {fmt(r['ci_high'])}", fmt(r["frac_negative"])]
           for r in rows_of("disagreement_definition_plddt_summary.csv")])

# S2: confound controls
cap("S2")
asym = rows_of("confound_control_asymmetry.csv")
mr_e = sorted(float(r["rho_evo_struct"]) for r in asym)
mr_s = sorted(float(r["rho_seq_struct"]) for r in asym)
med = lambda v: v[len(v) // 2]
para("(i) Asymmetry of the residue-level association:", bold=True, size=10)
add_table(["Pair", "median rho vs pLDDT", "n proteins"],
          [["evolution vs structure", f"{med(mr_e):.3f}", len(asym)],
           ["sequence vs structure", f"{med(mr_s):.3f}", len(asym)],
           ["% proteins with more negative evo-vs-struct rho",
            f"{sum(1 for r in asym if float(r['rho_evo_struct']) < float(r['rho_seq_struct'])) / len(asym) * 100:.0f}%", "-"]])
para("")
para("(ii) Structure-model DMS association, median |Spearman rho| with the DMS "
     "fitness score (released structure scores are oriented to deleteriousness, "
     "so signed correlations are negative):", bold=True, size=10)
add_table(["pLDDT bin", "model", "n variants", "n assays", "median |rho| vs DMS"],
          [[r["plddt_bin"], r["model"], fmt(r["n_variants"]), r["n_assays"],
            fmt(abs(float(r["median_rho_dms"])))]
           for r in rows_of("confound_control_accuracy.csv")])
para("")
para("(iii) Structure-model score coverage:", bold=True, size=10)
add_table(["pLDDT bin", "model", "fraction scored"],
          [[r["plddt_bin"], r["model"], fmt(r["frac_scored"])]
           for r in rows_of("confound_control_coverage.csv")])
para("")
para("(iv) Protein-fixed-effects regression of residue-level "
     "evo-vs-structure disagreement on pLDDT (dependent variable: "
     "per-residue median disagreement; within-protein demeaning = protein "
     "dummies; protein-clustered standard errors; beta per 10 pLDDT units):",
     bold=True, size=10)
mv = rows_of("multivariable_regression.csv")
add_table(["Model", "beta pLDDT (per 10 units)", "95% CI", "clustered SE",
           "P", "n residues", "proteins", "R2"],
          [[r["model"], f"{float(r['beta_plddt']) * 10:.4f}",
            f"[{float(r['ci_low']) * 10:.4f}, {float(r['ci_high']) * 10:.4f}]",
            f"{float(r['se_plddt']) * 10:.4f}", f"{float(r['p_plddt']):.4f}",
            fmt(r["n"]), r["n_proteins"], f"{float(r['r2']):.4f}"]
           for r in mv])
para("Adjusted model covariates (standardized direction of median "
     "coefficients): disorder beta -0.046 (P 0.082), contact density +0.006, "
     "helix -0.008, sheet -0.033, normalized position, MSA-depth category and "
     "structure coverage (full coefficients in "
     "results/statistics/multivariable_regression.csv). The pLDDT term remains "
     "negative and significant after adjustment.", size=9)

# S3: regime robustness across K
cap("S3")
ari = {(r["K_other"], r["K_ref"]): float(r["ARI"])
       for r in rows_of("regime_k_ari.csv")}
data3 = []
for r in rows_of("regime_k_robustness.csv"):
    k = r["K"]
    a = "1.000" if k == "6" else fmt(ari.get((k, "6")))
    data3.append([k, r["has_structure_dissenting"], fmt(r["consensus_tolerant_Y"]),
                  fmt(r["consensus_damaging_Y"]), fmt(r["mean_crossassay_agreement"]),
                  r["n_crossassay_pairs"], a])
add_table(["K", "structure-dissenting regime present", "Y consensus-tolerant",
           "Y consensus-damaging", "cross-assay agreement", "cross-assay pairs",
           "ARI vs K=6"], data3)

# S4: standard homology sensitivity
cap("S4")
para("Best-hit pairwise identity and query coverage from HMMER3 phmmer "
     "(clinical targets queried against the 217 DMS assay targets, "
     "E <= 1e-3); exclusion requires both criteria. Median AUROC uses the "
     "rank-averaged uniform ensemble over core-panel models.", size=9)
add_table(["Identity threshold", "query coverage", "proteins excluded",
           "shared with old 35", "proteins retained (with AUROC)",
           "median uniform AUROC"],
          [[f"pident >= {r['threshold_pident']}%",
            f">= {float(r['qcov_min_pct']):.0f}%",
            r["n_excluded"], r["overlap_with_old35"],
            fmt(r["n_proteins_with_auroc"]), fmt(r["median_uniform_auroc"])]
           for r in rows_of("homology_standard_sensitivity.csv")])

# S5: XGBoost ablations
cap("S5")
add_table(["Configuration", "median protein-level Spearman"],
          [[r["config"], fmt(r["median_rho"])]
           for r in rows_of("xgboost_ablation.csv")])

# S6: model panel composition
cap("S6")
add_table(["Model", "Official family", "Analysis family", "Uses MSA",
           "Uses structure", "Structure input AlphaFold", "In mechanistic panel",
           "Notes"],
          [[r["model"], r["official_family"], r["family_in_analysis"],
            r["uses_MSA"], r["uses_structure"], r["structure_input_is_AlphaFold"],
            r["in_mechanistic_panel"], r.get("notes", "")]
           for r in rows_of("model_panel_table.csv")])
para("Official family = model_type from the ProteinGym v1.3 config (the source "
     "of all modality flags). Tranception/TranceptEVE use an MSA only for "
     "inference-time homolog retrieval (not structure input); MIF-ST transfers "
     "sequence representations from a pretrained single-sequence language model "
     "(no MSA input). All structure-based scores in the released ProteinGym "
     "evaluation operate on the AlphaFold2 target structures shipped with the "
     "benchmark.", size=9)

# S7: clinical paired deltas
cap("S7")
add_table(["Set", "Comparison", "n proteins", "median delta AUROC",
           "95% CI", "frac positive"],
          [[r["set"], r["comparison"], fmt(r["n_proteins"]), fmt(r["median_delta"]),
            f"[{fmt(r['ci_low'])}, {fmt(r['ci_high'])}]", fmt(r["frac_positive"])]
           for r in rows_of("clinical_paired_deltas.csv")])

# S8: held-out regime generalization
cap("S8")
h = rows_of("heldout_regime_summary.csv")[0]
add_table(["Statistic", "Value"],
          [["Proteins with held-out support", h["n_proteins"]],
           ["Median Spearman, held-out regime profile vs training profile",
            f"{fmt(h['median_rho_train_profile'])} (IQR {fmt(h['q25_rho'])}-{fmt(h['q75_rho'])})"],
           ["Fraction of proteins with positive correlation",
            f"{float(h['frac_rho_positive']):.3f} ({h['n_proteins']}/{h['n_proteins']})"],
           ["Dissenting regime held-out Y (median)", fmt(h["dissenting_Y_median_heldout"])],
           ["Fraction of supported proteins with dissenting regime Y > 0.5",
            f"{float(h['frac_dissenting_Y_gt05']):.3f}"],
           ["Held-out cross-assay replication, median rho",
            f"{fmt(h['heldout_crossassay_rho_median'])} ({h['heldout_crossassay_n_pairs']} pairs)"],
           ["Fraction of held-out cross-assay pairs above chance",
            f"{float(h['frac_pairs_above_chance']):.3f} "
            f"({h['heldout_crossassay_n_pairs']}/{h['heldout_crossassay_n_pairs']})"]])
para("Per-protein values: results/statistics/heldout_regime_per_protein.csv.",
     size=9)

# S10: matched five-model control + averaging-convention sensitivity
cap("S10")
para("(a) Identical SD(z) disagreement and rank-averaged ensemble error in "
     "both datasets (five released zero-shot models; PoET oriented per "
     "official clinical metadata); median per-assay (DMS) / per-protein "
     "(clinical) Spearman with 10,000-resample bootstrap CIs:", bold=True,
     size=10)
m5 = rows_of("matched5_dms_vs_clinical.csv")
add_table(["Dataset", "statistic", "units", "median", "95% CI"],
          [[r["dataset"], r["statistic"], r["units"], fmt(r["value"]),
            f"{fmt(r['ci_low'])} to {fmt(r['ci_high'])}"] for r in m5])
para("")
para("(b) Five-model uniform AUROC by averaging convention:", bold=True,
     size=10)
add_table(["Set", "n proteins", "naive raw-score mean", "rank mean"],
          [[r["set"], r["n"], fmt(r["median_raw_mean"]),
            fmt(r["median_rank_mean"])]
           for r in rows_of("uniform_2x2_clinical.csv")])
para("")
para("(c) Paired BioGate minus rank-averaged five-model uniform (strict "
     "gate-complete set):", bold=True, size=10)
add_table(["Comparison", "n", "median delta", "95% CI", "frac positive"],
          [[r["comparison"], r["n"], fmt(r["median_delta"]),
            f"{fmt(r['ci_low'])} to {fmt(r['ci_high'])}",
            fmt(r["frac_positive"])]
           for r in rows_of("gate_vs_rankuniform.csv")])

# S11: structure-subset sensitivity
cap("S11")
add_table(["Structure configuration", "n models", "proteins", "median rho",
           "bootstrap 95% CI", "frac negative", "sign-test P"],
          [[r["config"], r["n_struct_models"], r["n_proteins"],
            fmt(r["median_rho"]), f"{fmt(r['ci_low'])} to {fmt(r['ci_high'])}",
            fmt(r["frac_negative"]), f"{float(r['sign_p']):.1e}"]
           for r in rows_of("struct_subset_sensitivity.csv")])

# S12: held-out without consensus regimes
cap("S12")
para("(a) Concordance of held-out regime phenotypes with training profiles "
     "(n >= 100 variants per held-out protein; IQR shown because the reduced-"
     "regime Spearman is discrete on 4 regimes):", bold=True, size=10)
add_table(["Metric", "value", "IQR", "proteins"],
          [[r["metric"], fmt(r["value"]),
            f"{fmt(r['ci_low'])} to {fmt(r['ci_high'])}"
            if r["ci_low"] not in ("", "nan") else "-",
            r["n"]]
           for r in rows_of("heldout_noconsensus_summary.csv")])
para("")
para("(b) Per-regime experimental-Y medians, training vs held-out (pooled "
     "across folds):", bold=True, size=10)
add_table(["Regime (machine label)", "held-out variants", "train Y median",
           "held-out Y median", "abs. difference"],
          [[f"regime_{int(float(r['reg']))}", fmt(r["heldout_n"]),
            fmt(r["train_Y"]), fmt(r["heldout_Y"]),
            f"{abs(float(r['train_Y']) - float(r['heldout_Y'])):.3f}"]
           for r in rows_of("heldout_noconsensus_regimes.csv")])

# S9: BIC sweep + functional-feature permutation enrichment
cap("S9")
para("(a) GMM model selection (family-centroid input, covariance_type=full, "
     "n_init=3, identical settings to the regime fit):", bold=True, size=10)
add_table(["K", "BIC", "selected"],
          [[r["K"], f"{float(r['BIC']):,.0f}", "yes" if r["selected"] == "True" else ""]
           for r in rows_of("gmm_bic_table.csv")])
para("")
para("(b) Within-protein permutation enrichment of high-disagreement residues "
     "for UniProt functional features (BH FDR across features):",
     bold=True, size=10)
add_table(["Feature", "proteins", "mean OR (high vs low decile)",
           "mean permutation P", "BH q"],
          [[r["feature"], r["n_proteins"] if r["n_proteins"] != "0" else "-",
            r["mean_or"] if r["mean_or"] != "inf" else "undefined",
            fmt(r["mean_perm_p"]), fmt(r["BH_q"])]
           for r in rows_of("functional_enrichment.csv")])
para("No feature survives BH control (minimum q = 0.60); disagreement "
     "structure is not a proxy for functional-site annotation.", size=9)

doc.save(str(OUT / "supplementary.docx"))
print("supplementary.docx:", (OUT / "supplementary.docx").stat().st_size,
      "bytes | tables:", len(TBL_CAP), "| figs:", len(FIG_S))
